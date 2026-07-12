from io import BytesIO

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.db.models import Base, SyncRun, Trade, User
from app.services.analytics_service import dashboard_data
from app.services.statement_import_service import (
    StatementFormatError,
    import_angel_pnl_statement,
    parse_angel_pnl_statement,
)


def statement_workbook() -> bytes:
    workbook = Workbook()
    equity = workbook.active
    equity.title = "Equity"
    equity["B7"] = "Client ID"
    equity["C7"] = "CLIENT01"
    equity["B11"] = "P&L Statement for Equity from 2026-01-01 to 2026-07-12"
    equity["B15"] = "Charges"
    equity["C15"] = 10
    equity["B16"] = "Other Credit & Debit"
    equity["C16"] = -15
    equity["B17"] = "Realized P&L"
    equity["C17"] = 50
    equity["B18"] = "Unrealized P&L"
    equity["C18"] = 20
    equity["B23"] = "Account Head"
    equity["C23"] = "Amount"
    equity["B24"] = "Brokerage"
    equity["C24"] = 10
    headers = [
        "Symbol",
        "ISIN",
        "Quantity",
        "Buy Value",
        "Sell Value",
        "Realized P&L",
        "Realized P&L Pct.",
        "Previous Closing Price",
        "Open Quantity",
        "Open Quantity Type",
        "Open Value",
        "Unrealized P&L",
        "Unrealized P&L Pct.",
    ]
    for column, value in enumerate(headers, start=2):
        equity.cell(row=39, column=column, value=value)
    realized = ["ABC", "INE000000001", 10, 1000, 1050, 50, 5, 0, 0, "", 0, 0, 0]
    opened = ["XYZ", "INE000000002", 0, 0, 0, 0, 0, 120, 5, "", 580, 20, 3.4483]
    for row_index, values in ((40, realized), (41, opened)):
        for column, value in enumerate(values, start=2):
            equity.cell(row=row_index, column=column, value=value)

    adjustments = workbook.create_sheet("Other Debits and Credits")
    adjustment_headers = ["Particulars", "Posting Date", "Debit", "Credit"]
    for column, value in enumerate(adjustment_headers, start=2):
        adjustments.cell(row=9, column=column, value=value)
    adjustment = ["DP charge", "2026-07-10", 15, 0]
    for column, value in enumerate(adjustment, start=2):
        adjustments.cell(row=10, column=column, value=value)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_statement_parser_reads_realized_open_and_adjustment_data():
    statement = parse_angel_pnl_statement(statement_workbook())

    assert statement.client_id == "CLIENT01"
    assert statement.period_end.isoformat() == "2026-07-12"
    assert statement.summary["Realized P&L"] == 50
    assert statement.charges == [{"account_head": "Brokerage", "amount": 10}]
    assert statement.adjustments[0]["debit"] == 15
    assert len(statement.equity_rows) == 2


def test_statement_import_replaces_previous_statement_records():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        user = User(email="statement@tradeos.app", name="Statement", password_hash="unused")
        db.add(user)
        db.commit()
        db.refresh(user)
        statement = parse_angel_pnl_statement(statement_workbook())

        first = import_angel_pnl_statement(db, user.id, statement, "statement.xlsx")
        second = import_angel_pnl_statement(db, user.id, statement, "statement.xlsx")

        assert first["realized_positions"] == 1
        assert first["open_positions"] == 1
        assert second["records_imported"] == 4
        assert db.scalar(select(func.count(Trade.id))) == 2
        assert db.scalar(select(func.count(SyncRun.id))) == 2
        metrics = dashboard_data(db, user.id)["metrics"]
        assert metrics["trade_count"] == 1
        assert metrics["net_pnl"] == 25


def test_statement_parser_rejects_non_xlsx_content():
    with pytest.raises(StatementFormatError, match="valid XLSX"):
        parse_angel_pnl_statement(b"not an Excel workbook")
