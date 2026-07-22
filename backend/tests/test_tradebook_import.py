import re
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.db.models import Base, SyncRun, Trade, TradeExecution, User
from app.services.analytics_service import dashboard_data
from app.services.tradebook_import_service import (
    TradebookFormatError,
    import_angel_tradebook,
    parse_angel_tradebook,
)

HEADERS = [
    "Symbol",
    "ISIN",
    "Trade Date",
    "Exchange",
    "Segment",
    "Series",
    "Trade Type",
    "Auction",
    "Quantity",
    "Price",
    "Trade ID",
    "Order ID",
    "Order Execution Time",
]


def execution(
    symbol: str,
    isin: str,
    trade_date: str,
    side: str,
    quantity: int,
    price: float,
    trade_id: str,
    order_id: str,
    executed_at: str,
) -> list:
    return [
        symbol,
        isin,
        trade_date,
        "NSE",
        "EQ",
        "EQ",
        side,
        False,
        quantity,
        price,
        trade_id,
        order_id,
        executed_at,
    ]


BASE_EXECUTIONS = [
    execution("ABC", "INE000000001", "2026-01-02", "buy", 6, 100, "B1", "ORDER1", "2026-01-02T09:15:00"),
    execution("ABC", "INE000000001", "2026-01-02", "buy", 4, 102, "B2", "ORDER1", "2026-01-02T09:16:00"),
    execution("XYZ", "INE000000002", "2026-01-02", "sell", 3, 50, "X1", "ORDER2", "2026-01-02T09:30:00"),
    execution("ABC", "INE000000001", "2026-01-03", "sell", 7, 110, "S1", "ORDER3", "2026-01-03T10:00:00"),
]


def tradebook_workbook(rows: list[list], malformed_dimension: bool = False) -> bytes:
    workbook = Workbook()
    equity = workbook.active
    equity.title = "Equity"
    equity["B7"] = "Client ID"
    equity["C7"] = "CLIENT01"
    equity["B11"] = "Tradebook for Equity from 2026-01-01 to 2026-01-31"
    for column, value in enumerate(HEADERS, start=2):
        equity.cell(row=15, column=column, value=value)
    for row_index, values in enumerate(rows, start=16):
        for column, value in enumerate(values, start=2):
            equity.cell(row=row_index, column=column, value=value)
    output = BytesIO()
    workbook.save(output)
    content = output.getvalue()
    if not malformed_dimension:
        return content

    rewritten = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(rewritten, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(br'<dimension ref="[^"]+"', b'<dimension ref="A1"', data, count=1)
            target.writestr(member, data)
    return rewritten.getvalue()


def test_tradebook_parser_handles_angel_one_malformed_dimension_and_split_orders():
    tradebook = parse_angel_tradebook(tradebook_workbook(BASE_EXECUTIONS, malformed_dimension=True))

    assert tradebook.client_id == "CLIENT01"
    assert tradebook.period_start.isoformat() == "2026-01-02"
    assert tradebook.period_end.isoformat() == "2026-01-03"
    assert len(tradebook.executions) == 4
    assert len({item.execution_key for item in tradebook.executions}) == 4
    assert len({item.order_id for item in tradebook.executions}) == 3


def test_tradebook_import_is_idempotent_and_adds_only_incremental_executions():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        user = User(email="tradebook@tradeos.app", name="Tradebook", password_hash="unused")
        db.add(user)
        db.commit()
        db.refresh(user)

        initial = parse_angel_tradebook(tradebook_workbook(BASE_EXECUTIONS))
        first = import_angel_tradebook(db, user.id, initial, "tradebook.xlsx")
        repeated = import_angel_tradebook(db, user.id, initial, "tradebook.xlsx")

        assert first["executions_added"] == 4
        assert first["closed_trades"] == 2
        assert first["open_lots"] == 1
        assert first["unmatched_sell_executions"] == 1
        assert first["unmatched_sell_quantity"] == 3
        assert repeated["executions_added"] == 0
        assert repeated["executions_existing"] == 4
        assert db.scalar(select(func.count(TradeExecution.id))) == 4
        assert db.scalar(select(func.count(Trade.id))) == 3

        incremental_rows = BASE_EXECUTIONS + [
            execution(
                "ABC",
                "INE000000001",
                "2026-01-04",
                "sell",
                3,
                90,
                "S2",
                "ORDER4",
                "2026-01-04T11:00:00",
            )
        ]
        incremental = parse_angel_tradebook(tradebook_workbook(incremental_rows))
        result = import_angel_tradebook(db, user.id, incremental, "tradebook-latest.xlsx")

        assert result["executions_added"] == 1
        assert result["executions_existing"] == 4
        assert result["total_executions"] == 5
        assert result["closed_trades"] == 3
        assert result["open_lots"] == 0
        assert db.scalar(select(func.count(TradeExecution.id))) == 5
        assert db.scalar(select(func.count(Trade.id))) == 3
        assert db.scalar(select(func.count(SyncRun.id))) == 3

        trades = list(db.scalars(select(Trade).order_by(Trade.exit_time, Trade.entry_time)).all())
        assert [trade.trade_date.isoformat() for trade in trades] == ["2026-01-03", "2026-01-03", "2026-01-04"]
        assert [trade.quantity for trade in trades] == [6, 1, 3]
        assert sum(trade.net_pnl for trade in trades) == pytest.approx(32)
        assert all(trade.raw_data["match_method"] == "FIFO" for trade in trades)
        assert dashboard_data(db, user.id)["metrics"]["net_pnl"] == 32

        sliced = parse_angel_tradebook(
            tradebook_workbook(
                [
                    execution(
                        "DEF",
                        "INE000000003",
                        "2026-01-05",
                        "buy",
                        2,
                        200,
                        "D1",
                        "ORDER5",
                        "2026-01-05T09:45:00",
                    )
                ]
            )
        )
        sliced_result = import_angel_tradebook(db, user.id, sliced, "tradebook-daily.xlsx")

        assert sliced_result["executions_added"] == 1
        assert sliced_result["executions_existing"] == 0
        assert sliced_result["total_executions"] == 6
        assert sliced_result["period_start"] == "2026-01-02"
        assert sliced_result["period_end"] == "2026-01-05"
        assert db.scalar(select(func.count(TradeExecution.id))) == 6
        assert db.scalar(select(func.count(SyncRun.id))) == 4


def test_tradebook_import_rejects_changed_values_for_an_existing_execution():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        user = User(email="conflict@tradeos.app", name="Conflict", password_hash="unused")
        db.add(user)
        db.commit()
        db.refresh(user)
        initial = parse_angel_tradebook(tradebook_workbook([BASE_EXECUTIONS[0]]))
        import_angel_tradebook(db, user.id, initial, "tradebook.xlsx")
        changed = parse_angel_tradebook(
            tradebook_workbook(
                [
                    execution(
                        "ABC",
                        "INE000000001",
                        "2026-01-02",
                        "buy",
                        6,
                        101,
                        "B1",
                        "ORDER1",
                        "2026-01-02T09:15:00",
                    )
                ]
            )
        )

        with pytest.raises(TradebookFormatError, match="conflicts with a previously imported execution"):
            import_angel_tradebook(db, user.id, changed, "tradebook-corrected.xlsx")


def test_tradebook_import_aggregates_remaining_buy_lots_into_one_open_position():
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        user = User(email="positions@tradeos.app", name="Positions", password_hash="unused")
        db.add(user)
        db.commit()
        db.refresh(user)
        tradebook = parse_angel_tradebook(tradebook_workbook(BASE_EXECUTIONS[:2]))

        result = import_angel_tradebook(db, user.id, tradebook, "tradebook.xlsx")

        assert result["open_lots"] == 2
        assert result["open_positions"] == 1
        opened = db.scalar(select(Trade).where(Trade.status == "OPEN"))
        assert opened is not None
        assert opened.quantity == 10
        assert opened.entry_price == pytest.approx(100.8)
        assert opened.product == "EQUITY"
        assert len(opened.raw_data["remaining_buy_lots"]) == 2


def test_tradebook_parser_rejects_non_xlsx_content():
    with pytest.raises(TradebookFormatError, match="valid XLSX"):
        parse_angel_tradebook(b"not an Excel workbook")


def test_tradebook_parser_rejects_a_pnl_statement():
    workbook = Workbook()
    workbook.active.title = "Equity"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(TradebookFormatError, match="does not contain an Angel One equity tradebook"):
        parse_angel_tradebook(output.getvalue())
