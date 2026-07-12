import hashlib
import re
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.db.models import BrokerAccount, SyncRun, Trade

MAX_WORKBOOK_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_WORKBOOK_FILES = 1_000
STATEMENT_TRADE_PREFIX = "STMT:"


class StatementFormatError(ValueError):
    pass


@dataclass
class ParsedStatement:
    client_id: str
    period_start: date
    period_end: date
    summary: dict[str, float]
    charges: list[dict]
    adjustments: list[dict]
    equity_rows: list[dict]
    checksum: str


def _number(value, default: float = 0.0) -> float:
    try:
        return float(value if value not in (None, "") else default)
    except (TypeError, ValueError):
        return default


def _integer(value) -> int:
    return int(round(abs(_number(value))))


def _json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_WORKBOOK_FILES:
                raise StatementFormatError("The workbook contains too many embedded files")
            if sum(member.file_size for member in members) > MAX_WORKBOOK_UNCOMPRESSED_BYTES:
                raise StatementFormatError("The workbook expands beyond the supported size")
            for member in members:
                path = PurePosixPath(member.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise StatementFormatError("The workbook contains an unsafe archive path")
            if "[Content_Types].xml" not in archive.namelist() or "xl/workbook.xml" not in archive.namelist():
                raise StatementFormatError("The uploaded file is not a valid XLSX workbook")
    except BadZipFile as exc:
        raise StatementFormatError("The uploaded file is not a valid XLSX workbook") from exc


def _sheet_rows(worksheet) -> list[list]:
    return [[_json_value(cell.value) for cell in row] for row in worksheet.iter_rows()]


def _find_header(rows: list[list], required: set[str]) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(rows):
        headers = {str(value).strip(): column for column, value in enumerate(row) if value not in (None, "")}
        if required.issubset(headers):
            return row_index, headers
    raise StatementFormatError(f"Missing required columns: {', '.join(sorted(required))}")


def _extract_period(rows: list[list]) -> tuple[date, date]:
    pattern = re.compile(r"from\s+(\d{4}-\d{2}-\d{2})\s+to\s+(\d{4}-\d{2}-\d{2})", re.IGNORECASE)
    for row in rows:
        for value in row:
            match = pattern.search(str(value or ""))
            if match:
                return date.fromisoformat(match.group(1)), date.fromisoformat(match.group(2))
    raise StatementFormatError("The statement reporting period could not be found")


def _extract_client_id(rows: list[list]) -> str:
    for row in rows:
        for index, value in enumerate(row[:-1]):
            if str(value or "").strip().lower() == "client id":
                client_id = str(row[index + 1] or "").strip()
                if client_id:
                    return client_id
    raise StatementFormatError("The Angel One client ID could not be found")


def _extract_summary(rows: list[list]) -> dict[str, float]:
    names = {"Charges", "Other Credit & Debit", "Realized P&L", "Unrealized P&L"}
    summary: dict[str, float] = {}
    for row in rows:
        for index, value in enumerate(row[:-1]):
            label = str(value or "").strip()
            if label in names and row[index + 1] not in (None, "") and label not in summary:
                summary[label] = _number(row[index + 1])
    if "Realized P&L" not in summary:
        raise StatementFormatError("The statement summary is incomplete")
    return summary


def _extract_equity_rows(rows: list[list]) -> list[dict]:
    required = {"Symbol", "ISIN", "Quantity", "Buy Value", "Sell Value", "Realized P&L", "Open Quantity"}
    header_index, headers = _find_header(rows, required)
    result = []
    for row in rows[header_index + 1 :]:
        symbol = str(row[headers["Symbol"]] or "").strip() if len(row) > headers["Symbol"] else ""
        if not symbol:
            continue
        result.append({header: row[column] if column < len(row) else None for header, column in headers.items()})
    if not result:
        raise StatementFormatError("The statement contains no equity rows")
    return result


def _extract_charges(rows: list[list]) -> list[dict]:
    header_index, headers = _find_header(rows, {"Account Head", "Amount"})
    result = []
    for row in rows[header_index + 1 :]:
        label = str(row[headers["Account Head"]] or "").strip() if len(row) > headers["Account Head"] else ""
        if not label or label == "Symbol":
            break
        result.append({"account_head": label, "amount": _number(row[headers["Amount"]])})
    return result


def _extract_adjustments(rows: list[list]) -> list[dict]:
    header_index, headers = _find_header(rows, {"Particulars", "Posting Date", "Debit", "Credit"})
    result = []
    for row in rows[header_index + 1 :]:
        particulars = str(row[headers["Particulars"]] or "").strip() if len(row) > headers["Particulars"] else ""
        if not particulars:
            continue
        posting_value = row[headers["Posting Date"]]
        try:
            posting_date = date.fromisoformat(str(posting_value)[:10]).isoformat()
        except ValueError:
            posting_date = str(posting_value or "")
        result.append(
            {
                "particulars": particulars,
                "posting_date": posting_date,
                "debit": _number(row[headers["Debit"]]),
                "credit": _number(row[headers["Credit"]]),
            }
        )
    return result


def parse_angel_pnl_statement(content: bytes) -> ParsedStatement:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True, keep_links=False)
    except Exception as exc:
        raise StatementFormatError("The XLSX workbook could not be read") from exc
    if "Equity" not in workbook.sheetnames or "Other Debits and Credits" not in workbook.sheetnames:
        raise StatementFormatError("Expected the Equity and Other Debits and Credits worksheets")

    equity_rows = _sheet_rows(workbook["Equity"])
    adjustment_rows = _sheet_rows(workbook["Other Debits and Credits"])
    period_start, period_end = _extract_period(equity_rows)
    return ParsedStatement(
        client_id=_extract_client_id(equity_rows),
        period_start=period_start,
        period_end=period_end,
        summary=_extract_summary(equity_rows),
        charges=_extract_charges(equity_rows),
        adjustments=_extract_adjustments(adjustment_rows),
        equity_rows=_extract_equity_rows(equity_rows),
        checksum=hashlib.sha256(content).hexdigest(),
    )


def _normalized_trades(statement: ParsedStatement) -> list[dict]:
    realized_rows = [
        row
        for row in statement.equity_rows
        if _integer(row.get("Quantity")) or _number(row.get("Buy Value")) or _number(row.get("Sell Value"))
    ]
    statement_costs = _number(statement.summary.get("Charges")) - _number(statement.summary.get("Other Credit & Debit"))
    turnovers = [_number(row.get("Buy Value")) + _number(row.get("Sell Value")) for row in realized_rows]
    total_turnover = sum(turnovers)
    remaining_costs = statement_costs
    trades: list[dict] = []

    for index, row in enumerate(realized_rows):
        quantity = _integer(row.get("Quantity"))
        if index == len(realized_rows) - 1:
            allocated_costs = remaining_costs
        elif total_turnover:
            allocated_costs = statement_costs * turnovers[index] / total_turnover
            remaining_costs -= allocated_costs
        else:
            allocated_costs = statement_costs / len(realized_rows)
            remaining_costs -= allocated_costs
        gross_pnl = _number(row.get("Realized P&L"))
        trades.append(
            {
                "broker_trade_id": (
                    f"{STATEMENT_TRADE_PREFIX}{statement.period_start}:{statement.period_end}:{row.get('ISIN')}:REALIZED"
                ),
                "symbol": str(row.get("Symbol")),
                "exchange": "NSE",
                "segment": "Equity",
                "product": "DELIVERY",
                "direction": "LONG",
                "quantity": quantity,
                "entry_price": _number(row.get("Buy Value")) / quantity if quantity else 0,
                "exit_price": _number(row.get("Sell Value")) / quantity if quantity else None,
                "gross_pnl": gross_pnl,
                "charges": allocated_costs,
                "net_pnl": gross_pnl - allocated_costs,
                "return_percent": _number(row.get("Realized P&L Pct.")),
                "status": "CLOSED",
                "trade_date": statement.period_end,
                "raw_data": {
                    "source": "angel_pnl_statement",
                    "date_basis": "statement_period_end",
                    "statement_row": row,
                    "allocated_statement_costs": allocated_costs,
                },
            }
        )

    for row in statement.equity_rows:
        quantity = _integer(row.get("Open Quantity"))
        if not quantity:
            continue
        open_value = _number(row.get("Open Value"))
        unrealized_pnl = _number(row.get("Unrealized P&L"))
        previous_close = _number(row.get("Previous Closing Price"))
        trades.append(
            {
                "broker_trade_id": (
                    f"{STATEMENT_TRADE_PREFIX}{statement.period_start}:{statement.period_end}:{row.get('ISIN')}:OPEN"
                ),
                "symbol": str(row.get("Symbol")),
                "exchange": "NSE",
                "segment": "Equity",
                "product": "DELIVERY",
                "direction": "LONG",
                "quantity": quantity,
                "entry_price": open_value / quantity if quantity else 0,
                "exit_price": previous_close or None,
                "gross_pnl": unrealized_pnl,
                "charges": 0,
                "net_pnl": unrealized_pnl,
                "return_percent": _number(row.get("Unrealized P&L Pct.")),
                "status": "OPEN",
                "trade_date": statement.period_end,
                "raw_data": {
                    "source": "angel_pnl_statement",
                    "date_basis": "statement_period_end",
                    "statement_row": row,
                },
            }
        )
    return trades


def import_angel_pnl_statement(db: Session, user_id: int, statement: ParsedStatement, filename: str) -> dict:
    account = db.scalar(
        select(BrokerAccount)
        .where(BrokerAccount.user_id == user_id, BrokerAccount.mode == "statement")
        .order_by(BrokerAccount.id.desc())
    )
    if account is None:
        account = BrokerAccount(user_id=user_id, mode="statement")
        db.add(account)
        db.flush()

    imported_at = datetime.now(UTC).replace(tzinfo=None)
    trades = _normalized_trades(statement)
    db.execute(delete(Trade).where(Trade.user_id == user_id, Trade.broker_account_id == account.id))
    for row in trades:
        db.add(
            Trade(
                user_id=user_id,
                broker_account_id=account.id,
                broker_trade_id=row["broker_trade_id"],
                symbol=row["symbol"],
                exchange=row["exchange"],
                segment=row["segment"],
                product=row["product"],
                direction=row["direction"],
                quantity=row["quantity"],
                entry_price=row["entry_price"],
                exit_price=row["exit_price"],
                gross_pnl=row["gross_pnl"],
                charges=row["charges"],
                net_pnl=row["net_pnl"],
                return_percent=row["return_percent"],
                status=row["status"],
                trade_date=row["trade_date"],
                holding_minutes=0,
                raw_data=row["raw_data"],
            )
        )

    realized_count = sum(trade["status"] == "CLOSED" for trade in trades)
    open_count = sum(trade["status"] == "OPEN" for trade in trades)
    account.broker_name = "Angel One Statement"
    account.mode = "statement"
    account.status = "connected"
    account.client_code = statement.client_id
    account.last_synced_at = imported_at
    account.raw_snapshot = {
        "profile": {"client_id": statement.client_id, "source": "pnl_statement"},
        "holdings": [row for row in statement.equity_rows if _integer(row.get("Open Quantity"))],
        "orders": [],
        "positions": [row for row in statement.equity_rows if _integer(row.get("Open Quantity"))],
        "rms": {},
        "statement": {
            "filename": filename,
            "checksum": statement.checksum,
            "period_start": statement.period_start.isoformat(),
            "period_end": statement.period_end.isoformat(),
            "summary": statement.summary,
            "charges": statement.charges,
            "adjustments": statement.adjustments,
            "equity_rows": statement.equity_rows,
        },
    }
    records = len(trades) + len(statement.charges) + len(statement.adjustments)
    sync_run = SyncRun(
        user_id=user_id,
        broker_account_id=account.id,
        status="completed",
        started_at=imported_at,
        completed_at=imported_at,
        records_synced=records,
        details={
            "source": "statement",
            "realized": realized_count,
            "open_positions": open_count,
            "charges": len(statement.charges),
            "adjustments": len(statement.adjustments),
        },
    )
    db.add(sync_run)
    db.commit()
    return {
        "status": "completed",
        "records_imported": records,
        "realized_positions": realized_count,
        "open_positions": open_count,
        "charges": len(statement.charges),
        "adjustments": len(statement.adjustments),
        "period_start": statement.period_start.isoformat(),
        "period_end": statement.period_end.isoformat(),
        "date_basis": "statement_period_end",
    }
