import hashlib
import re
from collections import defaultdict, deque
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.db.models import BrokerAccount, SyncRun, Trade, TradeExecution

MAX_WORKBOOK_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_WORKBOOK_FILES = 1_000
TRADEBOOK_TRADE_PREFIX = "TB:"
REQUIRED_HEADERS = {
    "symbol",
    "isin",
    "trade date",
    "exchange",
    "segment",
    "series",
    "trade type",
    "auction",
    "quantity",
    "price",
    "trade id",
    "order id",
    "order execution time",
}


class TradebookFormatError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedExecution:
    execution_key: str
    trade_id: str
    order_id: str
    symbol: str
    isin: str
    exchange: str
    segment: str
    series: str
    side: str
    auction: bool
    quantity: int
    price: float
    trade_date: date
    executed_at: datetime


@dataclass(frozen=True)
class ParsedTradebook:
    client_id: str
    period_start: date
    period_end: date
    executions: list[ParsedExecution]
    checksum: str


@dataclass
class FifoProjection:
    trades: list[dict]
    positions: list[dict]
    closed_trades: int
    open_lots: int
    unmatched_sell_executions: int
    unmatched_sell_quantity: int


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_WORKBOOK_FILES:
                raise TradebookFormatError("The workbook contains too many embedded files")
            if sum(member.file_size for member in members) > MAX_WORKBOOK_UNCOMPRESSED_BYTES:
                raise TradebookFormatError("The workbook expands beyond the supported size")
            for member in members:
                path = PurePosixPath(member.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise TradebookFormatError("The workbook contains an unsafe archive path")
            required_files = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required_files.issubset(archive.namelist()):
                raise TradebookFormatError("The uploaded file is not a valid XLSX workbook")
    except BadZipFile as exc:
        raise TradebookFormatError("The uploaded file is not a valid XLSX workbook") from exc


def _label(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def _identifier(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def _parse_date(value, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], pattern).date()
        except ValueError:
            continue
    raise TradebookFormatError(f"Row {row_number}: Trade Date is invalid")


def _parse_datetime(value, row_number: int) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        text = str(value or "").strip()
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError as exc:
            raise TradebookFormatError(f"Row {row_number}: Order Execution Time is invalid") from exc
    if parsed.tzinfo:
        return parsed.astimezone(UTC).replace(tzinfo=None)
    return parsed


def _parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return _label(value) in {"1", "true", "yes", "y"}


def _find_tradebook_sheet(workbook) -> tuple[object, list[tuple]]:
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        for row in rows[:100]:
            headers = {_label(value) for value in row if value not in (None, "")}
            if REQUIRED_HEADERS.issubset(headers):
                return worksheet, rows
    raise TradebookFormatError("The workbook does not contain an Angel One equity tradebook")


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(rows[:100]):
        headers = {_label(value): column for column, value in enumerate(row) if value not in (None, "")}
        if REQUIRED_HEADERS.issubset(headers):
            return row_index, headers
    raise TradebookFormatError("The tradebook is missing required execution columns")


def _extract_client_id(rows: list[tuple]) -> str:
    for row in rows[:100]:
        for index, value in enumerate(row[:-1]):
            if _label(value) == "client id":
                client_id = _identifier(row[index + 1])
                if client_id:
                    return client_id
    raise TradebookFormatError("The Angel One client ID could not be found")


def _parse_execution(row: tuple, headers: dict[str, int], row_number: int) -> ParsedExecution:
    def cell(name: str):
        column = headers[name]
        return row[column] if column < len(row) else None

    symbol = _identifier(cell("symbol")).upper()
    isin = _identifier(cell("isin")).upper()
    exchange = _identifier(cell("exchange")).upper()
    segment = _identifier(cell("segment")).upper()
    series = _identifier(cell("series")).upper()
    side = _label(cell("trade type")).upper()
    trade_id = _identifier(cell("trade id"))
    order_id = _identifier(cell("order id"))
    if not all((symbol, isin, exchange, segment, series, trade_id, order_id)):
        raise TradebookFormatError(f"Row {row_number}: required execution values are missing")
    if side not in {"BUY", "SELL"}:
        raise TradebookFormatError(f"Row {row_number}: Trade Type must be buy or sell")
    try:
        quantity_value = float(cell("quantity"))
        price = float(cell("price"))
    except (TypeError, ValueError) as exc:
        raise TradebookFormatError(f"Row {row_number}: Quantity or Price is invalid") from exc
    quantity = int(quantity_value)
    if quantity_value != quantity or quantity <= 0 or price <= 0:
        raise TradebookFormatError(f"Row {row_number}: Quantity and Price must be positive")
    trade_date = _parse_date(cell("trade date"), row_number)
    executed_at = _parse_datetime(cell("order execution time"), row_number)
    execution_key = f"{exchange}:{segment}:{trade_date.isoformat()}:{trade_id}"
    return ParsedExecution(
        execution_key=execution_key,
        trade_id=trade_id,
        order_id=order_id,
        symbol=symbol,
        isin=isin,
        exchange=exchange,
        segment=segment,
        series=series,
        side=side,
        auction=_parse_bool(cell("auction")),
        quantity=quantity,
        price=price,
        trade_date=trade_date,
        executed_at=executed_at,
    )


def parse_angel_tradebook(content: bytes) -> ParsedTradebook:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True, keep_links=False)
    except Exception as exc:
        raise TradebookFormatError("The XLSX workbook could not be read") from exc
    try:
        _, rows = _find_tradebook_sheet(workbook)
        header_index, headers = _find_header(rows)
        executions_by_key: dict[str, ParsedExecution] = {}
        for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            if not any(value not in (None, "") for value in row):
                continue
            if not _identifier(row[headers["symbol"]] if headers["symbol"] < len(row) else None):
                continue
            execution = _parse_execution(row, headers, row_index)
            existing = executions_by_key.get(execution.execution_key)
            if existing and existing != execution:
                raise TradebookFormatError(f"Row {row_index}: duplicate Trade ID has conflicting values")
            executions_by_key[execution.execution_key] = execution
        executions = sorted(
            executions_by_key.values(),
            key=lambda item: (item.executed_at, item.exchange, item.trade_id),
        )
        if not executions:
            raise TradebookFormatError("The tradebook contains no execution records")
        return ParsedTradebook(
            client_id=_extract_client_id(rows),
            period_start=min(item.trade_date for item in executions),
            period_end=max(item.trade_date for item in executions),
            executions=executions,
            checksum=hashlib.sha256(content).hexdigest(),
        )
    finally:
        workbook.close()


def _money(value: Decimal) -> float:
    return float(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _derived_trade_id(kind: str, *parts: str) -> str:
    digest = hashlib.sha256("|".join(parts).encode("utf-8")).hexdigest()
    return f"{TRADEBOOK_TRADE_PREFIX}{kind}:{digest}"


def _fifo_projection(executions: list[TradeExecution], namespace: str) -> FifoProjection:
    buys_by_isin: dict[str, deque[list]] = defaultdict(deque)
    trades: list[dict] = []
    unmatched_sell_executions = 0
    unmatched_sell_quantity = 0

    for execution in executions:
        if execution.side == "BUY":
            buys_by_isin[execution.isin].append([execution, execution.quantity])
            continue

        remaining = execution.quantity
        buy_lots = buys_by_isin[execution.isin]
        while remaining and buy_lots:
            buy, available = buy_lots[0]
            matched_quantity = min(remaining, available)
            cost = Decimal(str(buy.price)) * matched_quantity
            proceeds = Decimal(str(execution.price)) * matched_quantity
            gross_pnl = proceeds - cost
            return_percent = (gross_pnl / cost * 100) if cost else Decimal("0")
            holding_minutes = max(int((execution.executed_at - buy.executed_at).total_seconds() // 60), 0)
            trades.append(
                {
                    "broker_trade_id": _derived_trade_id(
                        "C", namespace, buy.execution_key, execution.execution_key, str(matched_quantity)
                    ),
                    "symbol": execution.symbol,
                    "exchange": (
                        execution.exchange
                        if buy.exchange == execution.exchange
                        else f"{buy.exchange}/{execution.exchange}"
                    ),
                    "segment": execution.segment,
                    "product": "EQUITY",
                    "direction": "LONG",
                    "quantity": matched_quantity,
                    "entry_price": buy.price,
                    "exit_price": execution.price,
                    "gross_pnl": _money(gross_pnl),
                    "charges": 0.0,
                    "net_pnl": _money(gross_pnl),
                    "return_percent": float(return_percent.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)),
                    "status": "CLOSED",
                    "trade_date": execution.trade_date,
                    "entry_time": buy.executed_at,
                    "exit_time": execution.executed_at,
                    "holding_minutes": holding_minutes,
                    "raw_data": {
                        "source": "angel_tradebook",
                        "match_method": "FIFO",
                        "buy_execution_key": buy.execution_key,
                        "sell_execution_key": execution.execution_key,
                        "charges_basis": "unavailable_in_tradebook",
                    },
                }
            )
            remaining -= matched_quantity
            available -= matched_quantity
            if available:
                buy_lots[0][1] = available
            else:
                buy_lots.popleft()
        if remaining:
            unmatched_sell_executions += 1
            unmatched_sell_quantity += remaining

    positions_by_isin: dict[str, dict] = {}
    open_lot_count = 0
    for isin, lots in buys_by_isin.items():
        for buy, quantity in lots:
            open_lot_count += 1
            position = positions_by_isin.setdefault(
                isin,
                {
                    "symbol": buy.symbol,
                    "isin": isin,
                    "exchanges": set(),
                    "segment": buy.segment,
                    "quantity": 0,
                    "cost": Decimal("0"),
                    "entry_time": buy.executed_at,
                    "trade_date": buy.trade_date,
                    "lots": [],
                },
            )
            position["exchanges"].add(buy.exchange)
            position["quantity"] += quantity
            position["cost"] += Decimal(str(buy.price)) * quantity
            position["entry_time"] = min(position["entry_time"], buy.executed_at)
            position["trade_date"] = min(position["trade_date"], buy.trade_date)
            position["lots"].append({"execution_key": buy.execution_key, "quantity": quantity})

    positions = []
    for isin, position in positions_by_isin.items():
        quantity = position["quantity"]
        average_price = _money(position["cost"] / quantity)
        exchange = "/".join(sorted(position["exchanges"]))
        lot_identity = [
            f"{lot['execution_key']}:{lot['quantity']}"
            for lot in sorted(position["lots"], key=lambda item: item["execution_key"])
        ]
        trades.append(
            {
                "broker_trade_id": _derived_trade_id("O", namespace, isin, *lot_identity),
                "symbol": position["symbol"],
                "exchange": exchange,
                "segment": position["segment"],
                "product": "EQUITY",
                "direction": "LONG",
                "quantity": quantity,
                "entry_price": average_price,
                "exit_price": None,
                "gross_pnl": 0.0,
                "charges": 0.0,
                "net_pnl": 0.0,
                "return_percent": 0.0,
                "status": "OPEN",
                "trade_date": position["trade_date"],
                "entry_time": position["entry_time"],
                "exit_time": None,
                "holding_minutes": 0,
                "raw_data": {
                    "source": "angel_tradebook",
                    "match_method": "FIFO",
                    "remaining_buy_lots": position["lots"],
                    "charges_basis": "unavailable_in_tradebook",
                },
            }
        )
        positions.append(
            {
                "symbol": position["symbol"],
                "isin": isin,
                "exchange": exchange,
                "quantity": quantity,
                "average_price": average_price,
            }
        )
    positions.sort(key=lambda item: item["symbol"])
    return FifoProjection(
        trades=trades,
        positions=positions,
        closed_trades=sum(trade["status"] == "CLOSED" for trade in trades),
        open_lots=open_lot_count,
        unmatched_sell_executions=unmatched_sell_executions,
        unmatched_sell_quantity=unmatched_sell_quantity,
    )


def _replace_derived_trades(
    db: Session,
    user_id: int,
    account_id: int,
    projection: FifoProjection,
) -> None:
    db.execute(delete(Trade).where(Trade.user_id == user_id, Trade.broker_account_id == account_id))
    for row in projection.trades:
        db.add(Trade(user_id=user_id, broker_account_id=account_id, **row))


def _same_execution(stored: TradeExecution, incoming: ParsedExecution) -> bool:
    return (
        stored.trade_id,
        stored.order_id,
        stored.symbol,
        stored.isin,
        stored.exchange,
        stored.segment,
        stored.series,
        stored.side,
        stored.auction,
        stored.quantity,
        Decimal(str(stored.price)),
        stored.trade_date,
        stored.executed_at,
    ) == (
        incoming.trade_id,
        incoming.order_id,
        incoming.symbol,
        incoming.isin,
        incoming.exchange,
        incoming.segment,
        incoming.series,
        incoming.side,
        incoming.auction,
        incoming.quantity,
        Decimal(str(incoming.price)),
        incoming.trade_date,
        incoming.executed_at,
    )


def import_angel_tradebook(
    db: Session,
    user_id: int,
    tradebook: ParsedTradebook,
    filename: str,
) -> dict:
    account = db.scalar(
        select(BrokerAccount)
        .where(
            BrokerAccount.user_id == user_id,
            BrokerAccount.mode == "tradebook",
            BrokerAccount.client_code == tradebook.client_id,
        )
        .order_by(BrokerAccount.id.desc())
    )
    if account is None:
        account = BrokerAccount(
            user_id=user_id,
            broker_name="Angel One Tradebook",
            mode="tradebook",
            status="connected",
            client_code=tradebook.client_id,
        )
        db.add(account)
        db.flush()

    execution_keys = [item.execution_key for item in tradebook.executions]
    existing_by_key = {
        item.execution_key: item
        for item in db.scalars(
            select(TradeExecution).where(
                TradeExecution.broker_account_id == account.id,
                TradeExecution.execution_key.in_(execution_keys),
            )
        ).all()
    }
    imported_at = datetime.now(UTC).replace(tzinfo=None)
    for item in tradebook.executions:
        existing = existing_by_key.get(item.execution_key)
        if existing:
            if not _same_execution(existing, item):
                raise TradebookFormatError(
                    f"Trade ID {item.trade_id} conflicts with a previously imported execution"
                )
            continue
        db.add(
            TradeExecution(
                user_id=user_id,
                broker_account_id=account.id,
                execution_key=item.execution_key,
                trade_id=item.trade_id,
                order_id=item.order_id,
                symbol=item.symbol,
                isin=item.isin,
                exchange=item.exchange,
                segment=item.segment,
                series=item.series,
                side=item.side,
                auction=item.auction,
                quantity=item.quantity,
                price=item.price,
                trade_date=item.trade_date,
                executed_at=item.executed_at,
                imported_at=imported_at,
                raw_data={"source": "angel_tradebook", "filename": filename},
            )
        )
    db.flush()

    all_executions = list(
        db.scalars(
            select(TradeExecution)
            .where(TradeExecution.user_id == user_id, TradeExecution.broker_account_id == account.id)
            .order_by(TradeExecution.executed_at, TradeExecution.exchange, TradeExecution.trade_id)
        ).all()
    )
    projection = _fifo_projection(all_executions, str(account.id))
    _replace_derived_trades(db, user_id, account.id, projection)

    executions_added = len(tradebook.executions) - len(existing_by_key)
    coverage_start = min(item.trade_date for item in all_executions)
    coverage_end = max(item.trade_date for item in all_executions)
    unique_orders = {item.order_id for item in all_executions}
    account.broker_name = "Angel One Tradebook"
    account.mode = "tradebook"
    account.status = "connected"
    account.last_synced_at = imported_at
    account.raw_snapshot = {
        "profile": {"client_id": tradebook.client_id, "source": "tradebook"},
        "holdings": projection.positions,
        "positions": projection.positions,
        "orders": [{"order_id": order_id} for order_id in sorted(unique_orders)],
        "rms": {},
        "tradebook": {
            "filename": filename,
            "checksum": tradebook.checksum,
            "period_start": coverage_start.isoformat(),
            "period_end": coverage_end.isoformat(),
            "latest_file_period_start": tradebook.period_start.isoformat(),
            "latest_file_period_end": tradebook.period_end.isoformat(),
            "total_executions": len(all_executions),
            "closed_trades": projection.closed_trades,
            "open_lots": projection.open_lots,
            "open_positions": len(projection.positions),
            "unmatched_sell_executions": projection.unmatched_sell_executions,
            "unmatched_sell_quantity": projection.unmatched_sell_quantity,
        },
    }
    details = {
        "source": "tradebook",
        "executions_received": len(tradebook.executions),
        "executions_added": executions_added,
        "executions_existing": len(existing_by_key),
        "total_executions": len(all_executions),
        "closed_trades": projection.closed_trades,
        "open_lots": projection.open_lots,
        "open_positions": len(projection.positions),
        "unmatched_sell_executions": projection.unmatched_sell_executions,
        "unmatched_sell_quantity": projection.unmatched_sell_quantity,
    }
    db.add(
        SyncRun(
            user_id=user_id,
            broker_account_id=account.id,
            status="completed",
            started_at=imported_at,
            completed_at=imported_at,
            records_synced=executions_added,
            details=details,
        )
    )
    db.commit()
    return {
        "status": "completed",
        **details,
        "period_start": coverage_start.isoformat(),
        "period_end": coverage_end.isoformat(),
        "open_positions": len(projection.positions),
        "charges_included": False,
    }
